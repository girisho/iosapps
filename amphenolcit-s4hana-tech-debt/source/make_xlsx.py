import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L
import sys

OUT = sys.argv[1]

DARK = "17313B"; TEAL = "0F766E"; TINT = "E6F2F0"; AMBER = "F59E0B"; GREY = "F3F4F6"
F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=DARK)
sub_fill = PatternFill("solid", fgColor=TEAL)
in_fill = PatternFill("solid", fgColor="FFF4CC")   # input cells
calc_fill = PatternFill("solid", fgColor=GREY)
body = Font(name=F, size=10)
blue = Font(name=F, size=10, color="0000FF")
bold = Font(name=F, size=10, bold=True)
title_font = Font(name=F, size=16, bold=True, color=DARK)
thin = Side(style="thin", color="D1D5DB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()


def header(ws, row, cols, fill=hdr_fill):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = hdr_font; cell.fill = fill; cell.alignment = center; cell.border = box


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[L(i)].width = w


def title(ws, text, subtitle):
    ws["A1"] = text; ws["A1"].font = title_font
    ws["A2"] = subtitle; ws["A2"].font = Font(name=F, size=10, italic=True, color="4B5563")


# ------------------------------------------------------------------ Lists
lists = wb.active
lists.title = "Lists"
LISTS = {
    "ItemType": ["Feature", "Defect", "Tech Debt", "Enabler", "Risk/Compliance"],
    "Dimension": ["Process", "Extensibility", "Data", "Integration", "Operations", "UX"],
    "Area": ["FI", "CO", "SD", "MM", "PP", "QM", "PM", "EWM/WM", "Basis", "BTP", "Integration", "Cross-app"],
    "Status": ["New", "Triaged", "Ready", "In Progress", "Testing", "Done", "Rejected", "Deferred"],
    "Quarter": ["Q4-2026", "Q1-2027", "Q2-2027", "Q3-2027", "Q4-2027", "Backlog"],
    "Fib": [1, 2, 3, 5, 8, 13, 20],
    "YesNo": ["Y", "N"],
    "CCLevel": ["A", "B", "C", "D", "n/a"],
    "Quadrant": ["Deliberate-Prudent", "Deliberate-Reckless", "Inadvertent-Prudent", "Inadvertent-Reckless", "n/a"],
    "Severity": ["P1", "P2", "P3", "P4", "n/a"],
}
title(lists, "Lists (drop-down sources)", "Edit these lists to change the drop-downs in the Tech Debt Register. Keep one value per cell.")
for ci, (name, vals) in enumerate(LISTS.items(), 1):
    c = lists.cell(row=4, column=ci, value=name); c.font = hdr_font; c.fill = hdr_fill; c.alignment = center
    for ri, v in enumerate(vals, 5):
        lists.cell(row=ri, column=ci, value=v).font = body
    lists.column_dimensions[L(ci)].width = 20


def list_ref(name):
    ci = list(LISTS).index(name) + 1
    n = len(LISTS[name])
    return f"Lists!${L(ci)}$5:${L(ci)}${4 + n}"


# ------------------------------------------------------------------ Settings
st = wb.create_sheet("Settings")
title(st, "Settings & Assumptions", "Yellow cells are inputs. All values below are placeholders - replace with AmphenolCIT actuals.")
settings = [
    ("Blended day rate (USD / person-day)", 750, "Placeholder: blend of internal + AMS partner rates. Replace with finance-approved rate."),
    ("Working days per quarter", 60, "Assumption: ~12 weeks x 5 days after holidays."),
    ("Team availability for backlog work (%)", 0.7, "Assumption: 30% reserved for run/ops tickets, meetings, leave."),
    ("Sprint length (working days)", 10, "2-week sprints."),
    ("Stale item threshold (days)", 90, "Items older than this in New/Triaged are flagged in the register."),
    ("Plan start date", dt.date(2026, 10, 5), "First sprint start (Monday)."),
]
header(st, 4, ["Setting", "Value", "Note / source"])
for i, (k, v, n) in enumerate(settings, 5):
    st.cell(row=i, column=1, value=k).font = body
    c = st.cell(row=i, column=2, value=v); c.font = blue; c.fill = in_fill; c.border = box
    if isinstance(v, float) and v < 1: c.number_format = "0%"
    elif isinstance(v, dt.date): c.number_format = "yyyy-mm-dd"
    elif k.startswith("Blended"): c.number_format = "$#,##0"
    st.cell(row=i, column=3, value=n).font = Font(name=F, size=9, italic=True, color="4B5563")
widths(st, [42, 16, 80])
RATE = "Settings!$B$5"; QDAYS = "Settings!$B$6"; AVAIL = "Settings!$B$7"; STALE = "Settings!$B$9"

# ------------------------------------------------------------------ Register
rg = wb.create_sheet("Tech Debt Register", 1)
title(rg, "Unified Backlog & Tech Debt Register",
      "One backlog for Features, Defects, Tech Debt, Enablers and Risk/Compliance. Yellow = input; grey = calculated. Rows EX-### are ILLUSTRATIVE - overwrite with real items.")
cols = ["ID", "Title", "Item Type", "Clean Core Dimension", "Area / Module", "Description / symptom", "Root cause (Fowler quadrant)",
        "Clean Core Level (A-D)", "Blocks innovation? (Y/N)", "Bug severity (P1-P4)",
        "Business Value", "Time Criticality", "Risk Reduction / Opportunity Enablement", "Job Size", "WSJF score", "Priority rank",
        "Interest (hrs / month lost)", "Principal (effort person-days)", "Est. cost (USD)", "Annual interest cost (USD)", "Payback (months)",
        "Business owner", "IT owner", "Status", "Target quarter", "Date raised", "Date closed", "Age (days)", "Stale?", "Notes / evidence link"]
HR = 4
header(rg, HR, cols)
rg.row_dimensions[HR].height = 48

S = [  # ID, title, type, dim, area, desc, quadrant, cc, block, sev, BV, TC, RR, size, interest, principal, bowner, iowner, status, quarter, raised, closed, notes
    ("EX-001", "Business data held in custom Z-fields on material master (MARA append)", "Tech Debt", "Data", "MM",
     "Z-fields not exposed via CDS/OData; standard Fiori apps, APIs and analytics cannot see the data. Users revert to SE16N/Excel.",
     "Deliberate-Prudent", "C", "Y", "n/a", 13, 8, 13, 8, 40, 25, "Materials Mgmt lead", "MM functional lead", "Triaged", "Q1-2027", dt.date(2026, 9, 1), None,
     "Evidence: SE11 append ZAMARA; map each Z-field to standard field or Custom Fields app"),
    ("EX-002", "Activate Fiori launchpad + role-based spaces for AP / AR clerks", "Feature", "UX", "FI",
     "Finance users on SAP GUI only; Manage Supplier Line Items, Clear GL Accounts etc. not in use.",
     "n/a", "A", "N", "n/a", 13, 5, 8, 5, 20, 15, "Finance controller", "Fiori/Basis lead", "Ready", "Q4-2026", dt.date(2026, 9, 5), None,
     "Use Fiori Apps Library recommendations report from ST03N usage"),
    ("EX-003", "SAP Build Work Zone (standard) subaccount, IdP and content federation", "Enabler", "UX", "BTP",
     "No central entry point; users bookmark GUI transactions and separate portals.",
     "n/a", "A", "Y", "n/a", 8, 5, 13, 5, 0, 20, "CIO office", "BTP platform lead", "Ready", "Q4-2026", dt.date(2026, 9, 5), None,
     "Prereq: SAP Cloud Identity Services, Cloud Connector, content provider from S/4 FLP"),
    ("EX-004", "Point-to-point RFC / file interfaces to MES re-platformed to Integration Suite", "Tech Debt", "Integration", "Integration",
     "12 custom RFC/file jobs, no monitoring, failures found by users.",
     "Inadvertent-Prudent", "B", "Y", "n/a", 8, 5, 13, 13, 30, 60, "Plant ops director", "Integration lead", "New", "Q2-2027", dt.date(2026, 9, 10), None,
     "Run Integration Suite Migration Assessment + ISA-M classification"),
    ("EX-005", "Intercompany STO pricing condition wrong for plant 2100", "Defect", "Process", "SD",
     "Wrong transfer price on intercompany billing; manual correction each month-end.",
     "Inadvertent-Reckless", "n/a", "N", "P3", 5, 8, 5, 2, 16, 3, "Sales ops", "SD functional lead", "In Progress", "Q4-2026", dt.date(2026, 8, 20), None, ""),
    ("EX-006", "Modifications in sales-order user exits (MV45AFZZ) refactored to released BAdIs", "Tech Debt", "Extensibility", "SD",
     "Implicit enhancements + modifications slow every upgrade and block clean core.",
     "Deliberate-Reckless", "D", "Y", "n/a", 5, 3, 13, 8, 12, 30, "Sales ops", "ABAP dev lead", "Triaged", "Q2-2027", dt.date(2026, 9, 12), None,
     "ATC clean-core check variant findings"),
    ("EX-007", "Decommission unused custom code (no usage in 12 months)", "Tech Debt", "Extensibility", "Cross-app",
     "Approx. 30% of Z-objects show no execution in SCMON/SUSG data.",
     "Inadvertent-Prudent", "C", "N", "n/a", 3, 2, 8, 3, 8, 10, "CIO office", "ABAP dev lead", "Triaged", "Q1-2027", dt.date(2026, 9, 12), None,
     "Custom Code Migration app scoping; retire before remediate"),
    ("EX-008", "Adopt Monitor Material Coverage / MRP Live Fiori apps for planners", "Feature", "Process", "PP",
     "Planners use MD04 + Excel; exception-based planning not used.",
     "n/a", "A", "N", "n/a", 13, 5, 5, 5, 25, 15, "Supply chain VP", "PP functional lead", "New", "Q1-2027", dt.date(2026, 9, 15), None, ""),
    ("EX-009", "S/4HANA Feature Pack / release currency plan", "Risk/Compliance", "Operations", "Basis",
     "System behind latest FPS; missing innovations and security fixes.",
     "Deliberate-Prudent", "n/a", "Y", "n/a", 8, 8, 13, 13, 0, 45, "CIO office", "Basis lead", "New", "Q2-2027", dt.date(2026, 9, 15), None,
     "Run SAP Readiness Check for the target release"),
    ("EX-010", "Duplicate business partners (customer/supplier) cleansing", "Tech Debt", "Data", "Cross-app",
     "Estimated duplicates cause blocked payments and credit errors.",
     "Inadvertent-Reckless", "n/a", "Y", "n/a", 8, 5, 8, 8, 30, 25, "Master data owner", "Data lead", "Triaged", "Q1-2027", dt.date(2026, 9, 1), None,
     "Profile with Data Quality rules; evaluate SAP MDG consolidation"),
    ("EX-011", "ORDERS IDoc failures from EDI partner", "Defect", "Integration", "SD",
     "Status 51 errors daily; orders re-keyed manually.",
     "Inadvertent-Prudent", "n/a", "N", "P2", 8, 13, 8, 2, 20, 4, "Customer service", "Integration lead", "In Progress", "Q4-2026", dt.date(2026, 9, 18), None, "BD87"),
    ("EX-012", "Integration Suite landing zone: naming, transport (cTMS), monitoring in Cloud ALM", "Enabler", "Integration", "BTP",
     "Integration Suite just started; no standards yet, risk of new debt.",
     "n/a", "A", "Y", "n/a", 5, 8, 13, 5, 0, 15, "CIO office", "Integration lead", "Ready", "Q4-2026", dt.date(2026, 9, 5), None, ""),
    ("EX-013", "Replace custom ALV reports with standard analytical Fiori apps / CDS", "Tech Debt", "UX", "FI",
     "40+ Z-reports duplicate standard analytics.",
     "Deliberate-Prudent", "C", "N", "n/a", 5, 2, 5, 8, 15, 30, "Finance controller", "ABAP dev lead", "New", "Q3-2027", dt.date(2026, 9, 20), None, ""),
    ("EX-014", "Work Zone site for plant supervisors (Fiori + non-SAP tiles)", "Feature", "UX", "PP",
     "Supervisors switch between 5 tools per shift.",
     "n/a", "A", "N", "n/a", 8, 3, 5, 5, 10, 15, "Plant ops director", "BTP platform lead", "New", "Q2-2027", dt.date(2026, 9, 20), None, "Depends on EX-003"),
    ("EX-015", "Move transport management to Cloud ALM Change & Deployment", "Enabler", "Operations", "Basis",
     "Manual transport sequencing causes import errors.",
     "Inadvertent-Prudent", "n/a", "N", "n/a", 3, 3, 8, 5, 12, 12, "CIO office", "Basis lead", "New", "Q2-2027", dt.date(2026, 9, 21), None, ""),
    ("EX-016", "Delivery note form prints wrong plant address", "Defect", "Process", "SD",
     "Cosmetic; workaround exists.",
     "Inadvertent-Prudent", "n/a", "N", "P4", 1, 1, 1, 1, 2, 1, "Shipping", "SD functional lead", "New", "Backlog", dt.date(2026, 6, 10), None, ""),
    ("EX-017", "Z-field duplicates standard field (sales region on customer)", "Tech Debt", "Data", "SD",
     "Same value kept in two places; reports disagree.",
     "Inadvertent-Reckless", "C", "Y", "n/a", 5, 3, 8, 3, 10, 8, "Sales ops", "Data lead", "Ready", "Q1-2027", dt.date(2026, 9, 1), None, ""),
    ("EX-018", "Automated regression tests for top 20 end-to-end processes", "Enabler", "Operations", "Cross-app",
     "Manual testing makes every change risky and slow.",
     "n/a", "n/a", "Y", "n/a", 8, 5, 13, 13, 40, 50, "CIO office", "QA lead", "New", "Q2-2027", dt.date(2026, 9, 21), None,
     "Cloud ALM test management + automation tool"),
    ("EX-019", "PI/PO interface inventory and migration wave plan", "Risk/Compliance", "Integration", "Integration",
     "Legacy middleware approaching end of mainstream maintenance.",
     "Deliberate-Prudent", "n/a", "Y", "n/a", 5, 8, 13, 5, 0, 10, "CIO office", "Integration lead", "Triaged", "Q1-2027", dt.date(2026, 9, 10), None,
     "Only if PI/PO is in landscape - confirm"),
    ("EX-020", "Joule / AI use-case readiness spike", "Feature", "Process", "Cross-app",
     "Business asking for AI features; need readiness view (licensing, clean core, data).",
     "n/a", "A", "N", "n/a", 5, 2, 8, 2, 0, 5, "CIO office", "Enterprise architect", "New", "Q3-2027", dt.date(2026, 9, 25), None, "Timeboxed spike"),
]
first = HR + 1
N_ROWS = 200
last = HR + N_ROWS
for r in range(first, last + 1):
    i = r - first
    row = S[i] if i < len(S) else None
    if row:
        (idv, t, typ, dim, area, desc, quad, cc, blk, sev, bv, tc, rr, size, interest, principal,
         bo, io, status, q, raised, closed, notes) = row
        vals = {1: idv, 2: t, 3: typ, 4: dim, 5: area, 6: desc, 7: quad, 8: cc, 9: blk, 10: sev,
                11: bv, 12: tc, 13: rr, 14: size, 17: interest, 18: principal, 22: bo, 23: io, 24: status,
                25: q, 26: raised, 27: closed, 30: notes}
        for c, v in vals.items():
            rg.cell(row=r, column=c, value=v)
    # formulas
    rg.cell(row=r, column=15, value=f'=IF(OR(N{r}="",N{r}=0),"",ROUND((K{r}+L{r}+M{r})/N{r},1))')
    rg.cell(row=r, column=16, value=f'=IF(OR(O{r}="",X{r}="Done",X{r}="Rejected"),"",COUNTIFS($O${first}:$O${last},">"&O{r},$X${first}:$X${last},"<>Done",$X${first}:$X${last},"<>Rejected")+1)')
    rg.cell(row=r, column=19, value=f'=IF(R{r}="","",R{r}*{RATE})')
    rg.cell(row=r, column=20, value=f'=IF(Q{r}="","",Q{r}*12*{RATE}/8)')
    rg.cell(row=r, column=21, value=f'=IF(OR(S{r}="",T{r}="",T{r}=0),"",ROUND(S{r}/(T{r}/12),1))')
    rg.cell(row=r, column=28, value=f'=IF(Z{r}="","",IF(AA{r}="",TODAY(),AA{r})-Z{r})')
    rg.cell(row=r, column=29, value=f'=IF(AB{r}="","",IF(AND(OR(X{r}="New",X{r}="Triaged"),AB{r}>{STALE}),"STALE",""))')
    for c in range(1, len(cols) + 1):
        cell = rg.cell(row=r, column=c)
        cell.border = box
        cell.alignment = wrap if c in (2, 6, 30) else Alignment(vertical="top", horizontal="center" if c not in (22, 23) else "left", wrap_text=True)
        if c in (15, 16, 19, 20, 21, 28, 29):
            cell.fill = calc_fill; cell.font = body
        else:
            cell.fill = in_fill if c not in (1,) else PatternFill(None)
            cell.font = blue if c in (11, 12, 13, 14, 17, 18) else body
        if c in (26, 27): cell.number_format = "yyyy-mm-dd"
        if c in (19, 20): cell.number_format = "$#,##0"
widths(rg, [9, 42, 14, 14, 11, 48, 20, 11, 11, 10, 10, 10, 13, 9, 9, 9, 11, 12, 12, 13, 10, 20, 20, 12, 11, 11, 11, 9, 9, 40])
rg.freeze_panes = "C5"
rg.auto_filter.ref = f"A{HR}:{L(len(cols))}{last}"

for colL, lname in [("C", "ItemType"), ("D", "Dimension"), ("E", "Area"), ("G", "Quadrant"), ("H", "CCLevel"), ("I", "YesNo"),
                    ("J", "Severity"), ("K", "Fib"), ("L", "Fib"), ("M", "Fib"), ("N", "Fib"), ("X", "Status"), ("Y", "Quarter")]:
    dv = DataValidation(type="list", formula1=f"={list_ref(lname)}", allow_blank=True)
    dv.error = "Pick a value from the list (edit the Lists tab to add values)."
    rg.add_data_validation(dv); dv.add(f"{colL}{first}:{colL}{last}")

# header comments = field help
helps = {
    "K": "Business Value (Fibonacci 1-20): revenue, cost, user productivity, compliance value to the business.",
    "L": "Time Criticality (1-20): does value decay if we wait? Fixed dates, month-end, audit, contract, maintenance deadlines.",
    "M": "Risk Reduction / Opportunity Enablement (1-20): does it reduce risk (security, upgrade, audit) or unlock other items (e.g. Fiori, Work Zone, AI)? Tech debt usually scores high here.",
    "N": "Job Size (1-20): relative effort. Use the same Fibonacci scale; NOT person-days.",
    "O": "WSJF = (Business Value + Time Criticality + Risk Reduction/Opportunity Enablement) / Job Size (SAFe).",
    "Q": "Interest: hours per month the organisation loses because the debt exists (workarounds, manual fixes, extra testing, incidents).",
    "R": "Principal: estimated person-days to remove the debt.",
    "T": "Annual interest cost = hrs/month x 12 x day rate / 8.",
    "U": "Payback = one-off fix cost / monthly interest cost. < 12 months is a strong candidate.",
    "H": "SAP clean core extensibility level: A = released APIs / ABAP Cloud; B = classic APIs; C = internal/unreleased objects; D = modifications / not recommended.",
    "G": "Fowler technical-debt quadrant: deliberate vs inadvertent, prudent vs reckless.",
    "J": "Only for Defects. P1/P2 go to the expedite lane and bypass WSJF.",
}
for colL, txt in helps.items():
    rg[f"{colL}{HR}"].comment = Comment(txt, "EA")

# conditional formats
rg.conditional_formatting.add(f"AC{first}:AC{last}", CellIsRule(operator="equal", formula=['"STALE"'], fill=PatternFill("solid", fgColor="FECACA"), font=Font(name=F, bold=True, color="991B1B")))
rg.conditional_formatting.add(f"P{first}:P{last}", CellIsRule(operator="between", formula=["1", "5"], fill=PatternFill("solid", fgColor="BBF7D0"), font=Font(name=F, bold=True)))
rg.conditional_formatting.add(f"J{first}:J{last}", FormulaRule(formula=[f'OR(J{first}="P1",J{first}="P2")'], fill=PatternFill("solid", fgColor="FECACA")))
rg.conditional_formatting.add(f"C{first}:C{last}", FormulaRule(formula=[f'C{first}="Tech Debt"'], font=Font(name=F, bold=True, color=TEAL)))

# ------------------------------------------------------------------ Capacity & Budget
cb = wb.create_sheet("Capacity & Budget", 2)
title(cb, "Capacity Allocation & Budget Guardrails",
      "Fund a stable team, then split its capacity by work type. Yellow = inputs. Demand is pulled live from the register (Principal person-days, open items).")
TYPES = LISTS["ItemType"]
QTRS = LISTS["Quarter"][:5]
cb["A4"] = "1. Capacity by quarter"; cb["A4"].font = bold
header(cb, 5, ["Quarter", "Team FTE", "Working days", "Availability %", "Capacity (person-days)", "Budget cap (USD)", "Capacity cost (USD)", "Within budget?"])
fte = [8, 8, 9, 9, 9]
budget = [380000, 380000, 420000, 420000, 420000]
for i, q in enumerate(QTRS):
    r = 6 + i
    cb.cell(row=r, column=1, value=q).font = bold
    c = cb.cell(row=r, column=2, value=fte[i]); c.font = blue; c.fill = in_fill
    cb.cell(row=r, column=3, value=f"={QDAYS}")
    cb.cell(row=r, column=4, value=f"={AVAIL}").number_format = "0%"
    cb.cell(row=r, column=5, value=f"=ROUND(B{r}*C{r}*D{r},0)")
    c = cb.cell(row=r, column=6, value=budget[i]); c.font = blue; c.fill = in_fill; c.number_format = "$#,##0"
    cb.cell(row=r, column=7, value=f"=B{r}*C{r}*{RATE}").number_format = "$#,##0"
    cb.cell(row=r, column=8, value=f'=IF(G{r}<=F{r},"Yes","OVER by "&TEXT(G{r}-F{r},"$#,##0"))')
    for cc in range(1, 9): cb.cell(row=r, column=cc).border = box
    for cc in (3, 4, 5, 7, 8): cb.cell(row=r, column=cc).fill = calc_fill
cb["J5"] = "Note"; cb["J5"].font = bold
cb["J6"] = "FTE and budget caps are placeholders. Capacity cost = FTE x working days x day rate (Settings)."
cb["J6"].font = Font(name=F, size=9, italic=True)

cb["A13"] = "2. Allocation guardrails (% of capacity by work type) - must total 100%"; cb["A13"].font = bold
header(cb, 14, ["Quarter"] + TYPES + ["Total", "Check"])
alloc = [  # Feature, Defect, TechDebt, Enabler, Risk
    [0.30, 0.20, 0.25, 0.20, 0.05],
    [0.35, 0.15, 0.25, 0.20, 0.05],
    [0.40, 0.15, 0.25, 0.15, 0.05],
    [0.45, 0.15, 0.20, 0.15, 0.05],
    [0.50, 0.15, 0.20, 0.10, 0.05],
]
for i, q in enumerate(QTRS):
    r = 15 + i
    cb.cell(row=r, column=1, value=q).font = bold
    for j, v in enumerate(alloc[i]):
        c = cb.cell(row=r, column=2 + j, value=v); c.font = blue; c.fill = in_fill; c.number_format = "0%"; c.border = box
    cb.cell(row=r, column=7, value=f"=SUM(B{r}:F{r})").number_format = "0%"
    cb.cell(row=r, column=8, value=f'=IF(ROUND(G{r},4)=1,"OK","Must = 100%")')
    cb.cell(row=r, column=7).fill = calc_fill; cb.cell(row=r, column=8).fill = calc_fill
cb.conditional_formatting.add("H15:H19", CellIsRule(operator="notEqual", formula=['"OK"'], fill=PatternFill("solid", fgColor="FECACA")))
cb.conditional_formatting.add("H6:H10", CellIsRule(operator="notEqual", formula=['"Yes"'], fill=PatternFill("solid", fgColor="FECACA")))

cb["A22"] = "3. Planned capacity (person-days) vs open demand from register"; cb["A22"].font = bold
hdr3 = ["Quarter"]
for t in TYPES: hdr3 += [f"{t} - capacity", f"{t} - demand"]
hdr3 += ["Total capacity", "Total demand", "Gap (+ = spare)"]
header(cb, 23, hdr3)
cb.row_dimensions[23].height = 42
REG = "'Tech Debt Register'"
for i, q in enumerate(QTRS):
    r = 24 + i
    cb.cell(row=r, column=1, value=q).font = bold
    for j, t in enumerate(TYPES):
        capc = 2 + 2 * j
        cb.cell(row=r, column=capc, value=f"=ROUND($E{6 + i}*{L(2 + j)}{15 + i},0)")
        cb.cell(row=r, column=capc + 1, value=(
            f'=SUMIFS({REG}!$R${first}:$R${last},{REG}!$Y${first}:$Y${last},$A{r},{REG}!$C${first}:$C${last},"{t}",'
            f'{REG}!$X${first}:$X${last},"<>Done",{REG}!$X${first}:$X${last},"<>Rejected")'))
    tc = 2 + 2 * len(TYPES)
    cap_cells = ",".join(f"{L(2 + 2 * j)}{r}" for j in range(len(TYPES)))
    dem_cells = ",".join(f"{L(3 + 2 * j)}{r}" for j in range(len(TYPES)))
    cb.cell(row=r, column=tc, value=f"=SUM({cap_cells})")
    cb.cell(row=r, column=tc + 1, value=f"=SUM({dem_cells})")
    cb.cell(row=r, column=tc + 2, value=f"={L(tc)}{r}-{L(tc + 1)}{r}")
    for cc in range(1, tc + 3):
        cb.cell(row=r, column=cc).border = box
        if cc > 1: cb.cell(row=r, column=cc).fill = calc_fill
gapcol = L(2 + 2 * len(TYPES) + 2)
cb.conditional_formatting.add(f"{gapcol}24:{gapcol}28", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FECACA")))
for j in range(len(TYPES)):
    capL, demL = L(2 + 2 * j), L(3 + 2 * j)
    cb.conditional_formatting.add(f"{demL}24:{demL}28", FormulaRule(formula=[f"{demL}24>{capL}24"], fill=PatternFill("solid", fgColor="FDE68A")))
cb["A30"] = "Amber = demand for that work type exceeds its guardrail -> re-prioritise, move the quarter, or consciously flex the allocation at the monthly portfolio review."
cb["A30"].font = Font(name=F, size=9, italic=True)
cb["A31"] = "Rule of thumb: never let Tech Debt + Enabler drop below ~25% while clean-core KPIs are red; P1/P2 defects are funded from the Defect bucket first."
cb["A31"].font = Font(name=F, size=9, italic=True)
widths(cb, [12] + [13] * 15)

# ------------------------------------------------------------------ Roadmap (Gantt)
rm = wb.create_sheet("Roadmap", 3)
title(rm, "15-Month Roadmap (Oct 2026 - Dec 2027)", "Enter start / end dates; bars draw automatically. Dates are proposals for discussion.")
months = [dt.date(2026 + (9 + k) // 12, (9 + k) % 12 + 1, 1) for k in range(15)]
header(rm, 4, ["Workstream", "Activity", "Owner", "Start", "End"] + [m.strftime("%b-%y") for m in months])
R = [
    ("Governance", "Stand up CCoE / design authority, backlog tool, cadences", "EA + CIO office", dt.date(2026, 10, 5), dt.date(2026, 11, 13)),
    ("Governance", "Discovery & baseline (Readiness Check, Pathfinder, ATC, SCMON, interface + data inventory)", "EA", dt.date(2026, 10, 5), dt.date(2026, 12, 18)),
    ("Governance", "Monthly portfolio review + quarterly roadmap refresh", "CIO office", dt.date(2026, 11, 1), dt.date(2027, 12, 31)),
    ("UX / Fiori", "Fiori foundation: FLP, spaces & pages, roles, performance", "Basis / Fiori lead", dt.date(2026, 10, 19), dt.date(2027, 1, 29)),
    ("UX / Fiori", "Wave 1 Fiori apps: Finance (AP/AR/GL)", "Finance + Fiori lead", dt.date(2026, 12, 1), dt.date(2027, 3, 31)),
    ("UX / Fiori", "Wave 2 Fiori apps: Supply chain (MRP, inventory, purchasing)", "SCM + Fiori lead", dt.date(2027, 3, 1), dt.date(2027, 7, 30)),
    ("UX / Fiori", "Wave 3 Fiori apps: Sales, quality, maintenance", "Business leads", dt.date(2027, 7, 1), dt.date(2027, 11, 30)),
    ("Work Zone", "Work Zone standard: subaccount, IdP, content federation", "BTP lead", dt.date(2026, 11, 2), dt.date(2027, 1, 29)),
    ("Work Zone", "Pilot site (finance) then plant supervisor site", "BTP lead", dt.date(2027, 2, 1), dt.date(2027, 6, 30)),
    ("Integration", "Integration Suite landing zone + standards + monitoring", "Integration lead", dt.date(2026, 10, 19), dt.date(2027, 1, 15)),
    ("Integration", "Interface inventory, ISA-M classification, wave plan", "Integration lead", dt.date(2026, 11, 1), dt.date(2027, 1, 31)),
    ("Integration", "Migration waves 1-3 (P2P / PI-PO -> Integration Suite)", "Integration lead", dt.date(2027, 2, 1), dt.date(2027, 12, 17)),
    ("Data", "Custom field inventory + map to standard / Custom Fields app", "Data lead", dt.date(2026, 10, 19), dt.date(2027, 1, 29)),
    ("Data", "Data quality rules, owners, dashboard; BP de-duplication", "Data lead", dt.date(2026, 12, 1), dt.date(2027, 6, 30)),
    ("Data", "MDG / governance tool decision + implementation start", "Data lead", dt.date(2027, 4, 1), dt.date(2027, 12, 17)),
    ("Clean Core", "Retire unused custom code", "ABAP lead", dt.date(2026, 12, 1), dt.date(2027, 3, 31)),
    ("Clean Core", "Remediate Level C/D objects; new builds ABAP Cloud / BTP only", "ABAP lead", dt.date(2027, 2, 1), dt.date(2027, 12, 17)),
    ("Platform / Ops", "Cloud ALM: requirements, test mgmt, change & deployment, monitoring", "Basis lead", dt.date(2026, 10, 19), dt.date(2027, 3, 31)),
    ("Platform / Ops", "Test automation for top-20 processes", "QA lead", dt.date(2027, 1, 4), dt.date(2027, 6, 30)),
    ("Platform / Ops", "Release / FPS upgrade (target release TBC)", "Basis lead", dt.date(2027, 7, 1), dt.date(2027, 10, 29)),
    ("Innovation", "AI / Joule readiness spike + Build Process Automation pilots", "EA", dt.date(2027, 5, 3), dt.date(2027, 9, 30)),
]
for i, (ws_, act, own, s, e) in enumerate(R):
    r = 5 + i
    for c, v in enumerate([ws_, act, own, s, e], 1):
        cell = rm.cell(row=r, column=c, value=v); cell.border = box; cell.font = body
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if c in (4, 5): cell.number_format = "yyyy-mm-dd"; cell.fill = in_fill; cell.font = blue
    for k in range(len(months)):
        col = 6 + k
        mL = L(col)
        # month start in header row as date calc: use DATE of header month
        m = months[k]
        rm.cell(row=r, column=col, value=f'=IF(AND($D{r}<=EOMONTH(DATE({m.year},{m.month},1),0),$E{r}>=DATE({m.year},{m.month},1)),1,"")').font = Font(name=F, size=8, color="FFFFFF")
        rm.cell(row=r, column=col).border = box
    rm.row_dimensions[r].height = 30
last_r = 4 + len(R)
rng = f"F5:{L(5 + len(months))}{last_r}"
palette = {"Governance": "475569", "UX / Fiori": "0F766E", "Work Zone": "14B8A6", "Integration": "B45309",
           "Data": "7C3AED", "Clean Core": "BE123C", "Platform / Ops": "1D4ED8", "Innovation": "F59E0B"}
for wsname, color in palette.items():
    rm.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(F5=1,$A5="{wsname}")'], fill=PatternFill("solid", fgColor=color), font=Font(color=color)))
widths(rm, [15, 52, 20, 11, 11] + [6.5] * len(months))
rm.freeze_panes = "F5"

# ------------------------------------------------------------------ Cadence
cd = wb.create_sheet("Cadence", 4)
title(cd, "Operating Cadence - Grooming, Review & Repetition", "The rhythm that keeps the backlog healthy and progress visible.")
header(cd, 4, ["Ceremony", "Frequency", "Duration", "Participants", "Inputs", "Outputs / decisions", "Owner"])
C = [
    ("Stand-up (per team)", "Daily", "15 min", "Delivery team", "Board", "Blockers raised; WIP respected", "Scrum master"),
    ("Intake triage", "Weekly", "45 min", "Product owner, functional leads, AMS lead", "New requests, incidents, ATC/monitoring findings", "Classify type, dimension, severity; P1/P2 -> expedite; reject duplicates", "Product owner"),
    ("Backlog refinement (grooming)", "Weekly / bi-weekly", "60-90 min", "PO, leads, architects", "Top ~30 ranked items", "Sized, WSJF scored, Definition of Ready met", "Product owner"),
    ("Sprint planning", "Every 2 weeks", "2 h", "Delivery team, PO", "Ready items + capacity guardrails", "Sprint goal; mix per allocation %", "Scrum master"),
    ("Sprint review / demo", "Every 2 weeks", "1 h", "Team + business users", "Done items", "Accepted work; feedback into backlog", "Product owner"),
    ("Retrospective", "Every 2 weeks", "45 min", "Delivery team", "Flow metrics", "1-2 improvement actions", "Scrum master"),
    ("Architecture / design authority", "Bi-weekly", "60 min", "EA, BTP, integration, data, security", "New designs, clean core exceptions", "Approve extension pattern (A-D level), integration pattern (ISA-M)", "Enterprise architect"),
    ("Portfolio & capacity review", "Monthly", "60 min", "CIO, business process owners, PO, EA", "KPI dashboard, capacity vs demand", "Re-balance allocation %, approve epics, budget check", "CIO office"),
    ("Roadmap refresh / PI planning", "Quarterly", "0.5-1 day", "All leads + business", "Roadmap, KPI trends, SAP release news", "Next-quarter plan, re-ranked epics", "EA + CIO office"),
    ("Tech debt register health check", "Quarterly", "60 min", "EA, dev lead, PO", "Stale items, ATC trend, SCMON data", "Close stale items, re-score, retire obsolete code", "Enterprise architect"),
    ("SAP innovation scan", "Semi-annual", "Half day", "EA, leads, SAP CSP / partner", "Release notes, Pathfinder, Process Insights, Fiori app library", "New feature candidates added to backlog", "Enterprise architect"),
]
for i, row in enumerate(C):
    for c, v in enumerate(row, 1):
        cell = cd.cell(row=5 + i, column=c, value=v); cell.font = body; cell.alignment = wrap; cell.border = box
widths(cd, [30, 16, 11, 32, 34, 44, 18])

# ------------------------------------------------------------------ Sprint tracker
sp = wb.create_sheet("Sprint Tracker", 5)
title(sp, "Sprint Tracker & Flow Distribution", "Enter story points planned and completed per work type each sprint. Flow distribution shows where capacity really went.")
spc = ["Sprint", "Start", "End", "Capacity (pts)"] + [f"Planned - {t}" for t in TYPES] + [f"Done - {t}" for t in TYPES] + \
      ["Planned total", "Done total", "Say/Do %", "Tech Debt share of done %", "Defect share of done %"]
header(sp, 4, spc)
sp.row_dimensions[4].height = 42
for i in range(13):
    r = 5 + i
    sp.cell(row=r, column=1, value=f"S{i + 1:02d}")
    sp.cell(row=r, column=2, value=f"=Settings!$B$10+{14 * i}").number_format = "yyyy-mm-dd"
    sp.cell(row=r, column=3, value=f"=B{r}+11").number_format = "yyyy-mm-dd"
    for c in range(4, 15):
        cell = sp.cell(row=r, column=c); cell.fill = in_fill; cell.font = blue
    pt, dn = f"E{r}:I{r}", f"J{r}:N{r}"
    sp.cell(row=r, column=15, value=f"=SUM({pt})")
    sp.cell(row=r, column=16, value=f"=SUM({dn})")
    sp.cell(row=r, column=17, value=f'=IF(O{r}=0,"",P{r}/O{r})').number_format = "0%"
    sp.cell(row=r, column=18, value=f'=IF(P{r}=0,"",L{r}/P{r})').number_format = "0%"
    sp.cell(row=r, column=19, value=f'=IF(P{r}=0,"",K{r}/P{r})').number_format = "0%"
    for c in range(1, 20):
        sp.cell(row=r, column=c).border = box
        if c >= 15: sp.cell(row=r, column=c).fill = calc_fill
# one example row
ex = [40, 12, 8, 10, 8, 2, 10, 6, 9, 6, 2]
for j, v in enumerate(ex):
    sp.cell(row=5, column=4 + j, value=v)
sp["A19"] = "Row S01 contains example values showing the expected format - overwrite with actuals."
sp["A19"].font = Font(name=F, size=9, italic=True)
widths(sp, [8, 11, 11, 10] + [10] * 10 + [10, 10, 9, 12, 12])
sp.freeze_panes = "B5"

# ------------------------------------------------------------------ KPI dashboard
kd = wb.create_sheet("KPI Dashboard", 1)
title(kd, "Visible Progress - KPI Dashboard", "Backlog counts update live from the register. KPI baselines/targets are inputs to agree with leadership.")
kd["A4"] = "Backlog by type and status (live)"; kd["A4"].font = bold
statuses = ["New", "Triaged", "Ready", "In Progress", "Testing", "Done"]
header(kd, 5, ["Item type"] + statuses + ["Open total", "Open effort (p-days)", "Avg WSJF (open)"])
for i, t in enumerate(TYPES):
    r = 6 + i
    kd.cell(row=r, column=1, value=t).font = bold
    for j, s in enumerate(statuses):
        kd.cell(row=r, column=2 + j, value=f'=COUNTIFS({REG}!$C${first}:$C${last},$A{r},{REG}!$X${first}:$X${last},"{s}")')
    kd.cell(row=r, column=8, value=f"=SUM(B{r}:F{r})")
    kd.cell(row=r, column=9, value=f'=SUMIFS({REG}!$R${first}:$R${last},{REG}!$C${first}:$C${last},$A{r},{REG}!$X${first}:$X${last},"<>Done",{REG}!$X${first}:$X${last},"<>Rejected")')
    kd.cell(row=r, column=10, value=f'=IFERROR(AVERAGEIFS({REG}!$O${first}:$O${last},{REG}!$C${first}:$C${last},$A{r},{REG}!$X${first}:$X${last},"<>Done",{REG}!$X${first}:$X${last},"<>Rejected"),"")').number_format = "0.0"
    for c in range(1, 11): kd.cell(row=r, column=c).border = box
tr = 6 + len(TYPES)
kd.cell(row=tr, column=1, value="Total").font = bold
for c in range(2, 10):
    kd.cell(row=tr, column=c, value=f"=SUM({L(c)}6:{L(c)}{tr - 1})").font = bold
    kd.cell(row=tr, column=c).border = box

kd["A13"] = "Health signals (live)"; kd["A13"].font = bold
sig = [
    ("Stale items (New/Triaged > threshold)", f'=COUNTIF({REG}!$AC${first}:$AC${last},"STALE")'),
    ("Open items blocking innovation", f'=COUNTIFS({REG}!$I${first}:$I${last},"Y",{REG}!$X${first}:$X${last},"<>Done",{REG}!$X${first}:$X${last},"<>Rejected")'),
    ("Open Level C/D clean-core items", f'=COUNTIFS({REG}!$H${first}:$H${last},"C",{REG}!$X${first}:$X${last},"<>Done")+COUNTIFS({REG}!$H${first}:$H${last},"D",{REG}!$X${first}:$X${last},"<>Done")'),
    ("Open P1/P2 defects", f'=COUNTIFS({REG}!$J${first}:$J${last},"P1",{REG}!$X${first}:$X${last},"<>Done")+COUNTIFS({REG}!$J${first}:$J${last},"P2",{REG}!$X${first}:$X${last},"<>Done")'),
    ("Annual interest cost of open tech debt (USD)", f'=SUMIFS({REG}!$T${first}:$T${last},{REG}!$C${first}:$C${last},"Tech Debt",{REG}!$X${first}:$X${last},"<>Done")'),
    ("Open tech debt principal (USD)", f'=SUMIFS({REG}!$S${first}:$S${last},{REG}!$C${first}:$C${last},"Tech Debt",{REG}!$X${first}:$X${last},"<>Done")'),
]
for i, (k, f) in enumerate(sig):
    r = 14 + i
    kd.cell(row=r, column=1, value=k).font = body
    c = kd.cell(row=r, column=5, value=f); c.font = bold; c.border = box; c.fill = calc_fill
    if "USD" in k: c.number_format = "$#,##0"
    kd.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)

kd["A22"] = "Outcome KPIs (agree baseline + targets; update 'Current' monthly)"; kd["A22"].font = bold
header(kd, 23, ["KPI", "Dimension", "Unit", "Baseline", "Target Q2-2027", "Target Q4-2027", "Current", "Direction", "RAG", "Source / tool"])
K = [
    ("Share of active users launching Fiori apps monthly", "UX", "%", 0.0, 0.40, 0.75, None, "Up", "FLP usage / ST03N"),
    ("Fiori apps in productive use", "UX", "#", 0, 40, 120, None, "Up", "FLP usage analytics"),
    ("Users onboarded to Work Zone sites", "UX", "#", 0, 150, 600, None, "Up", "Work Zone / IAS"),
    ("Interfaces on Integration Suite (of total)", "Integration", "%", 0.05, 0.35, 0.70, None, "Up", "Interface inventory"),
    ("Interface failures per month", "Integration", "#", 120, 60, 20, None, "Down", "Cloud ALM Integration Monitoring"),
    ("Custom objects (Z) in use", "Extensibility", "#", 2500, 2000, 1600, None, "Down", "SCMON / Custom Code Migration app"),
    ("ATC clean-core findings Level C/D", "Extensibility", "#", 900, 700, 450, None, "Down", "ATC clean-core check variant"),
    ("Business-data custom fields mapped to standard or exposed", "Data", "%", 0.0, 0.50, 0.90, None, "Up", "Custom field inventory"),
    ("Master data quality score (rule pass rate)", "Data", "%", 0.72, 0.85, 0.95, None, "Up", "DQ rules / MDG / Information Steward"),
    ("Median lead time request -> production (days)", "Operations", "days", 45, 30, 20, None, "Down", "Backlog tool"),
    ("Tech debt share of delivered capacity", "Operations", "%", 0.05, 0.25, 0.20, None, "Hold", "Sprint Tracker"),
    ("Behind latest S/4 FPS / release (months)", "Operations", "months", 12, 12, 3, None, "Down", "System info / Readiness Check"),
]
for i, row in enumerate(K):
    r = 24 + i
    for c, v in enumerate(row, 1):
        cell = kd.cell(row=r, column=c, value=v); cell.font = body; cell.border = box
        if c in (4, 5, 6, 7):
            cell.fill = in_fill; cell.font = blue
            if row[2] == "%": cell.number_format = "0%"
    kd.cell(row=r, column=9, value=(
        f'=IF(G{r}="","Not measured",IF(H{r}="Up",IF(G{r}>=E{r},"G",IF(G{r}>D{r},"A","R")),'
        f'IF(H{r}="Down",IF(G{r}<=E{r},"G",IF(G{r}<D{r},"A","R")),IF(ABS(G{r}-E{r})<=0.05*MAX(ABS(E{r}),1),"G","A"))))'))
    kd.cell(row=r, column=9).border = box; kd.cell(row=r, column=9).alignment = center
rag = f"I24:I{23 + len(K)}"
kd.conditional_formatting.add(rag, CellIsRule(operator="equal", formula=['"G"'], fill=PatternFill("solid", fgColor="BBF7D0")))
kd.conditional_formatting.add(rag, CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", fgColor="FDE68A")))
kd.conditional_formatting.add(rag, CellIsRule(operator="equal", formula=['"R"'], fill=PatternFill("solid", fgColor="FECACA")))
kd.cell(row=25 + len(K), column=1, value="Baselines/targets are illustrative placeholders; RAG compares Current to the Q2-2027 target (G = met, A = improving vs baseline, R = not improving).").font = Font(name=F, size=9, italic=True)
widths(kd, [44, 13, 9, 10, 12, 12, 10, 10, 9, 32])

ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Backlog by type and status"; ch.y_axis.title = "Items"
data = Reference(kd, min_col=2, max_col=7, min_row=5, max_row=5 + len(TYPES))
cats = Reference(kd, min_col=1, min_row=6, max_row=5 + len(TYPES))
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.height = 7.5; ch.width = 16
kd.add_chart(ch, "L4")

# ------------------------------------------------------------------ Read Me
rd = wb.create_sheet("Read Me", 0)
title(rd, "AmphenolCIT S/4HANA - Tech Debt & Innovation Planner", "Companion to the strategy deck and quick reference guide. Draft v0.1 - Enterprise Architecture.")
rows = [
    ("How to use", ""),
    ("1. Settings", "Set day rate, working days, availability and plan start date."),
    ("2. Tech Debt Register", "Single backlog for ALL work types. Log every item with the quick reference guide fields. EX-### rows are illustrative - overwrite them."),
    ("3. Score", "Business Value, Time Criticality, Risk Reduction/Opportunity Enablement and Job Size on 1-2-3-5-8-13-20. WSJF and rank calculate automatically. P1/P2 defects bypass ranking (expedite lane)."),
    ("4. Capacity & Budget", "Enter FTE and budget caps; set % allocation per work type per quarter. Compare guardrail capacity to live demand from the register."),
    ("5. Roadmap", "Adjust start/end dates; Gantt bars redraw."),
    ("6. Sprint Tracker", "Record planned vs done points per type each sprint to show real flow distribution."),
    ("7. KPI Dashboard", "Live backlog counts and health signals; update 'Current' KPI values monthly for the portfolio review."),
    ("8. Cadence", "The ceremonies that keep grooming and repetition going."),
    ("", ""),
    ("Colour legend", ""),
    ("Yellow fill, blue text", "Input - edit these"),
    ("Grey fill", "Formula - do not overwrite"),
    ("Red highlight", "Needs attention (stale item, P1/P2, over budget, allocation not 100%)"),
    ("", ""),
    ("Definitions", ""),
    ("Tech Debt", "Something built or configured that works today but increases the cost or risk of future change (e.g. modifications, Z-fields holding standard data, P2P interfaces, missing tests)."),
    ("Defect", "Something that does not work as designed. Severity P1-P4."),
    ("Feature", "New or improved business capability (e.g. a Fiori app, Work Zone site, new S/4 functionality)."),
    ("Enabler", "Platform / architecture work that makes features possible (e.g. Work Zone subaccount, Integration Suite landing zone, test automation)."),
    ("Risk/Compliance", "Security, audit, maintenance-deadline or licence-driven work."),
    ("Interest vs Principal", "Interest = recurring cost of living with the debt (hrs/month). Principal = one-off cost to remove it. High interest + low principal = pay down first."),
]
for i, (a, b) in enumerate(rows, 4):
    ca = rd.cell(row=i, column=1, value=a); cb_ = rd.cell(row=i, column=2, value=b)
    ca.font = bold if b == "" or a[:1].isdigit() else body
    cb_.font = body; cb_.alignment = wrap
    if a in ("How to use", "Colour legend", "Definitions"):
        ca.font = Font(name=F, bold=True, color="FFFFFF"); ca.fill = sub_fill; rd.cell(row=i, column=2).fill = sub_fill
rd["A15"].fill = in_fill; rd["A15"].font = blue
rd["A16"].fill = calc_fill
rd["A17"].fill = PatternFill("solid", fgColor="FECACA")
widths(rd, [28, 110])

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("Lists",)
wb.move_sheet("Lists", offset=len(wb.sheetnames))
wb.move_sheet("Settings", offset=-(wb.sheetnames.index("Settings") - (len(wb.sheetnames) - 2)))
wb.save(OUT)
print("saved", wb.sheetnames)
